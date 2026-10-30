"""
make_price_overrides.py - turn the hand-verified price workbook into the venue dataset used by BOTH
the benchmark (run_scenario.py / batch_runner.py) and the web app (frontend/).

Usage (from any folder):
    python make_price_overrides.py                                   # price_data.xlsx + davao_venues.json next to this script
    python make_price_overrides.py price_data.xlsx davao_venues.json # explicit paths

Reads the single "Prices" sheet. Every row is a venue in the dataset; there is no other source of venues.
Writes (next to this script):
  benchmark_venues.json - THE dataset: only these venues can appear in a candidate pool, with their verified
                          price and the category from the workbook. run_scenario.py refuses to run without it.
  price_overrides.csv   - the verified prices in readable form (venue_id,name,per_person_php,room_night_php,source,date)
  venue_whitelist.csv   - the venue list in readable form (venue_id,name)
and copies benchmark_venues.json to ../frontend/src/data/venues.json so the app uses the identical file
(plus venues.meta.json with its data_version).

Rules:
  - Every row must match a venue (by the Google place_id inside its 'Open in Maps' link, so renaming is safe)
    and must have a numeric REAL price; otherwise the script stops and nothing is written.
  - 'Room rate per night' -> room_night_php (lodging, per room per night); everything else -> per_person_php
    (admission, day pass, 1-hour service, one meal, free = 0).
  - 'Category' in the workbook overrides the category from davao_venues.json (fix tags there, not in the JSON).
IMPORTANT: changing this file's output changes the benchmark. batch_runner.py refuses to resume a run that
was started on a different dataset; re-run EVERY model you compare under new tags.
"""
import csv
import datetime as dt
import hashlib
import json
import math
import os
import shutil
import sys
import urllib.parse

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_JSON = os.path.join(HERE, "benchmark_venues.json")
APP_JSON = os.path.join(HERE, "..", "frontend", "src", "data", "venues.json")
APP_META = os.path.join(HERE, "..", "frontend", "src", "data", "venues.meta.json")
INTERESTS = {"Food", "Nature", "Culture", "Entertainment", "Shopping", "Wellness", "Accommodation", "Sightseeing"}


def quality_score(v):
    """Rating weighted by popularity. Precomputed here so Python and the app sort on the exact same number."""
    return (v.get("average_rating") or 0) * math.log10((v.get("rating_count") or 0) + 10)


def as_float(x):
    return float(x) if x is not None else None


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "price_data.xlsx")
    venues_json = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "davao_venues.json")
    with open(venues_json, encoding="utf-8") as f:
        venues = {v["place_id"]: v for v in json.load(f)}
    ws = load_workbook(xlsx)["Prices"]
    col = {c.value: i for i, c in enumerate(ws[1])}

    out, errors, seen = [], [], set()
    for n, r in enumerate(ws.iter_rows(min_row=2), start=2):
        name = r[col["Venue"]].value
        if not name:
            continue
        link = r[col["Maps link"]].hyperlink
        pid = urllib.parse.parse_qs(urllib.parse.urlparse(link.target).query).get("query_place_id", [None])[0] if link else None
        v = venues.get(pid)
        if not v:
            errors.append(f"row {n} ({name}): no venue matches its Maps link")
            continue
        if v["venue_id"] in seen:
            errors.append(f"row {n} ({name}): {v['venue_id']} is listed twice")
            continue
        real = r[col["REAL price (PHP)"]].value
        try:
            amount = float(str(real).replace(",", ""))
        except ValueError:
            errors.append(f"row {n} ({name}): REAL price is missing or not a number ({real!r})")
            continue
        category = [c.strip() for c in str(r[col["Category"]].value or "").split(",") if c.strip()]
        unknown = [c for c in category if c not in INTERESTS]
        if not category or unknown:
            errors.append(f"row {n} ({name}): bad Category {r[col['Category']].value!r}")
            continue
        room = (r[col["What the real price is"]].value or "").strip() == "Room rate per night"
        if room != ("Accommodation" in category):
            errors.append(f"row {n} ({name}): room rates are only for Accommodation venues "
                          f"(category {category}, price is {'a room rate' if room else 'per person'})")
            continue
        date = r[col["Date checked"]].value
        if isinstance(date, dt.datetime):
            date = date.date().isoformat()
        seen.add(v["venue_id"])
        out.append({
            "venue_id": v["venue_id"],
            "place_id": v["place_id"],
            "name": v["name"],
            "name_key": (v["name"] or "").strip().lower(),     # used to drop chain-branch duplicates
            "category": category,
            "price_unit": "room_night" if room else "per_person",
            "price_php": amount,
            "price_source": r[col["Source (URL / call / visit)"]].value or "",
            "price_date": str(date or ""),
            "average_rating": as_float(v.get("average_rating")),
            "rating_count": v.get("rating_count"),
            "quality_score": quality_score(v),
            "operating_hours": v.get("operating_hours"),
            "latitude": as_float(v.get("latitude")),
            "longitude": as_float(v.get("longitude")),
            "address": v.get("address"),
        })

    if errors:
        print("Nothing written - fix these rows in the workbook first:")
        for e in errors:
            print("  ", e)
        raise SystemExit(1)

    unusable = [x["name"] for x in out if x["latitude"] is None or x["longitude"] is None or not x["operating_hours"]]
    if unusable:
        print(f"Warning: {len(unusable)} venue(s) have no coordinates or no opening hours and can never be picked:",
              unusable)

    with open(OUT_JSON, "w", encoding="utf-8", newline="\n") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)
    with open(os.path.join(HERE, "price_overrides.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["venue_id", "name", "per_person_php", "room_night_php", "source", "date"])
        for x in out:
            room = x["price_unit"] == "room_night"
            w.writerow([x["venue_id"], x["name"], "" if room else x["price_php"], x["price_php"] if room else "",
                        x["price_source"], x["price_date"]])
    with open(os.path.join(HERE, "venue_whitelist.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["venue_id", "name"])
        w.writerows([[x["venue_id"], x["name"]] for x in out])
    print(f"Wrote benchmark_venues.json, price_overrides.csv and venue_whitelist.csv with {len(out)} venue(s).")

    if os.path.isdir(os.path.dirname(APP_JSON)):
        shutil.copyfile(OUT_JSON, APP_JSON)
        with open(OUT_JSON, "rb") as f:
            version = hashlib.sha256(f.read()).hexdigest()[:12]      # same as run_scenario.DATA_VERSION
        with open(APP_META, "w", encoding="utf-8", newline="\n") as f:
            json.dump({"data_version": version}, f)
        print(f"Copied the dataset to {os.path.normpath(APP_JSON)} (data_version {version})")
    else:
        print(f"Warning: {os.path.normpath(os.path.dirname(APP_JSON))} not found - the app was NOT updated.")


if __name__ == "__main__":
    main()
